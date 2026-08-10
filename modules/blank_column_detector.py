"""
Feature 3 & 4: Blank Column Detection & Data Region Integrity Analysis
Scans the worksheet used range for blank columns between populated columns.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

import numpy as np
import openpyxl
import pandas as pd

from modules.file_handler import get_column_letter, get_used_range_info


@dataclass
class BlankColumnFinding:
    sheet_name: str
    col_index: int          # 0-based
    col_letter: str
    col_header: Optional[str]
    position: str           # "leading", "internal", "trailing"
    risk_level: str         # "CRITICAL", "HIGH", "MEDIUM", "LOW"
    description: str
    populated_before: List[str] = field(default_factory=list)  # col letters before
    populated_after: List[str] = field(default_factory=list)   # col letters after


@dataclass
class DataRegionReport:
    sheet_name: str
    total_columns: int
    populated_columns: List[str]       # letters
    blank_columns: List[str]           # letters
    split_points: List[str]            # blank col letters causing splits
    regions: List[Dict]                # {"start": "A", "end": "AX", "cols": [...]}
    affected_region: Optional[Dict]    # region after first internal blank
    internal_blank_count: int
    risk_level: str


def detect_blank_columns(
    wb: openpyxl.Workbook,
    sheet_name: str,
    df: pd.DataFrame,
) -> Tuple[List[BlankColumnFinding], DataRegionReport]:
    """
    Main entry point. Returns (findings, region_report).
    """
    used = get_used_range_info(wb, sheet_name)
    col_has_data = used["col_has_data"]
    headers = used["headers"]
    max_col = used["max_col"]

    # Build column liveness from DataFrame as secondary check
    # (more accurate for data rows, openpyxl used for structure)
    df_liveness = _df_column_liveness(df)

    # Merge: a column is blank only if BOTH openpyxl and df say so
    # (df may have fewer columns if header row differs)
    merged_liveness: List[bool] = []
    for i in range(max_col):
        opxl_live = col_has_data[i] if i < len(col_has_data) else False
        df_live = df_liveness[i] if i < len(df_liveness) else False
        merged_liveness.append(opxl_live or df_live)

    findings: List[BlankColumnFinding] = []

    # Find leading / internal / trailing blanks
    first_pop = _first_index(merged_liveness, True)
    last_pop = _last_index(merged_liveness, True)

    internal_blanks: List[int] = []
    for i, live in enumerate(merged_liveness):
        if not live:
            if first_pop is not None and last_pop is not None:
                if first_pop < i < last_pop:
                    internal_blanks.append(i)

    # Build populated regions (sequences separated by internal blanks)
    regions = _build_regions(merged_liveness, first_pop, last_pop)

    split_points = [get_column_letter(i) for i in internal_blanks]

    # Build findings list
    for idx in internal_blanks:
        letter = get_column_letter(idx)
        header = headers[idx] if idx < len(headers) else None

        # populated cols before/after this blank
        pop_before = [get_column_letter(i) for i in range(idx) if merged_liveness[i]]
        pop_after = [
            get_column_letter(i)
            for i in range(idx + 1, max_col)
            if merged_liveness[i]
        ]

        findings.append(
            BlankColumnFinding(
                sheet_name=sheet_name,
                col_index=idx,
                col_letter=letter,
                col_header=str(header) if header is not None else None,
                position="internal",
                risk_level="CRITICAL",
                description=(
                    f"Blank column {letter} splits the dataset between populated data. "
                    "Excel sort/filter operations may treat this as the data boundary, "
                    "causing downstream columns to remain stationary while earlier "
                    "columns move — silently mismatching rows."
                ),
                populated_before=pop_before[-5:],  # show last 5
                populated_after=pop_after[:5],
            )
        )

    # Determine affected region (everything after first internal blank)
    affected_region = None
    if internal_blanks and len(regions) > 1:
        # Region after first split
        after_regions = [r for r in regions if r["start_idx"] > internal_blanks[0]]
        if after_regions:
            all_affected = []
            for r in after_regions:
                all_affected.extend(r["cols"])
            affected_region = {
                "start": after_regions[0]["start"],
                "end": after_regions[-1]["end"],
                "cols": all_affected,
                "count": len(all_affected),
            }

    # Risk level
    if len(internal_blanks) >= 2:
        risk = "CRITICAL"
    elif len(internal_blanks) == 1:
        risk = "CRITICAL"
    else:
        risk = "LOW"

    report = DataRegionReport(
        sheet_name=sheet_name,
        total_columns=max_col,
        populated_columns=[
            get_column_letter(i) for i, live in enumerate(merged_liveness) if live
        ],
        blank_columns=[
            get_column_letter(i) for i, live in enumerate(merged_liveness) if not live
        ],
        split_points=split_points,
        regions=regions,
        affected_region=affected_region,
        internal_blank_count=len(internal_blanks),
        risk_level=risk,
    )

    return findings, report


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _df_column_liveness(df: pd.DataFrame) -> List[bool]:
    """Return per-column liveness based on data rows (row 1 onward)."""
    result = []
    for col in df.columns:
        series = df[col]
        has_value = series.notna().any() and (series.astype(str).str.strip() != "").any()
        result.append(bool(has_value))
    return result


def _first_index(lst: List[bool], value: bool) -> Optional[int]:
    for i, v in enumerate(lst):
        if v == value:
            return i
    return None


def _last_index(lst: List[bool], value: bool) -> Optional[int]:
    for i in range(len(lst) - 1, -1, -1):
        if lst[i] == value:
            return i
    return None


def _build_regions(
    liveness: List[bool],
    first_pop: Optional[int],
    last_pop: Optional[int],
) -> List[Dict]:
    """Build contiguous populated regions separated by internal blanks."""
    if first_pop is None:
        return []

    regions = []
    in_region = False
    region_start = None
    region_cols = []

    for i in range(first_pop, (last_pop or 0) + 1):
        if liveness[i]:
            if not in_region:
                in_region = True
                region_start = i
                region_cols = []
            region_cols.append(get_column_letter(i))
        else:
            if in_region:
                regions.append({
                    "start": get_column_letter(region_start),
                    "end": get_column_letter(i - 1),
                    "start_idx": region_start,
                    "end_idx": i - 1,
                    "cols": list(region_cols),
                    "count": len(region_cols),
                })
                in_region = False
                region_start = None
                region_cols = []

    if in_region and region_start is not None:
        regions.append({
            "start": get_column_letter(region_start),
            "end": get_column_letter(last_pop),
            "start_idx": region_start,
            "end_idx": last_pop,
            "cols": list(region_cols),
            "count": len(region_cols),
        })

    return regions
