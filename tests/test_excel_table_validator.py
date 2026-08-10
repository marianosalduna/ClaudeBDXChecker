"""
Unit tests for Excel Table validation (Feature 5).
"""
import unittest
from unittest.mock import MagicMock, patch

from modules.excel_table_validator import validate_excel_tables


def _make_wb_with_table(table_name="BDXData", ref="A1:N200"):
    """Build a mock workbook where the sheet has an Excel Table."""
    tbl = MagicMock()
    tbl.ref = ref
    tbl.displayName = table_name

    ws = MagicMock()
    ws.tables = {table_name: tbl}
    ws.sheet_state = "visible"

    wb = MagicMock()
    wb.__getitem__ = MagicMock(return_value=ws)
    return wb


def _make_wb_without_table():
    """Build a mock workbook where the sheet has no Excel Table."""
    ws = MagicMock()
    ws.tables = {}
    ws.sheet_state = "visible"

    wb = MagicMock()
    wb.__getitem__ = MagicMock(return_value=ws)
    return wb


class TestValidateExcelTables(unittest.TestCase):
    def test_sheet_with_table_returns_has_table_true(self):
        wb = _make_wb_with_table("MyTable", "A1:Z100")

        with patch("modules.excel_table_validator.openpyxl") as mock_opxl:
            from openpyxl.utils import range_boundaries
            result = validate_excel_tables(wb, "Sheet1")

        self.assertTrue(result.has_table)
        self.assertEqual(result.risk_level, "LOW")
        self.assertEqual(len(result.tables), 1)
        self.assertEqual(result.tables[0].table_name, "MyTable")

    def test_sheet_without_table_returns_has_table_false(self):
        wb = _make_wb_without_table()
        result = validate_excel_tables(wb, "Sheet1")

        self.assertFalse(result.has_table)
        self.assertEqual(result.risk_level, "MEDIUM")
        self.assertEqual(len(result.tables), 0)

    def test_recommendation_present(self):
        wb = _make_wb_without_table()
        result = validate_excel_tables(wb, "Sheet1")
        self.assertTrue(len(result.recommendation) > 0)
        self.assertIn("Ctrl+T", result.recommendation)

    def test_message_present_for_table(self):
        wb = _make_wb_with_table()
        result = validate_excel_tables(wb, "Sheet1")
        self.assertIn("Table", result.message)

    def test_message_present_for_no_table(self):
        wb = _make_wb_without_table()
        result = validate_excel_tables(wb, "Sheet1")
        self.assertIn("NOT stored", result.message)

    def test_multiple_tables(self):
        ws = MagicMock()
        t1 = MagicMock()
        t1.ref = "A1:D100"
        t1.displayName = "Table1"
        t2 = MagicMock()
        t2.ref = "F1:J100"
        t2.displayName = "Table2"
        ws.tables = {"Table1": t1, "Table2": t2}
        ws.sheet_state = "visible"
        wb = MagicMock()
        wb.__getitem__ = MagicMock(return_value=ws)

        result = validate_excel_tables(wb, "Sheet1")
        self.assertTrue(result.has_table)
        self.assertEqual(len(result.tables), 2)


if __name__ == "__main__":
    unittest.main()
